from __future__ import annotations

import io
import json
import subprocess
import sys
import tempfile
import unittest
from dataclasses import asdict
from pathlib import Path
from unittest import mock

from passbolt_review import analyze_files


def write_encrypted_xlsx(path: Path, document_password: str, credential_password: str) -> None:
    from msoffcrypto.format.ooxml import OOXMLFile
    from openpyxl import Workbook

    plaintext = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Accessi"
    sheet.append(["Titolo", "Utente", "Password", "URL"])
    sheet.append(
        ["Gestionale protetto", "utente-protetto", credential_password, "10.30.40.50"]
    )
    workbook.save(plaintext)
    workbook.close()
    plaintext.seek(0)
    try:
        with path.open("wb") as encrypted_output:
            OOXMLFile(plaintext).encrypt(document_password, encrypted_output)
    finally:
        plaintext.close()


class ReviewTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)
        (self.root / "Cliente Alfa").mkdir()
        (self.root / "Cliente Beta").mkdir()

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_text_candidate_masks_and_never_serializes_secret(self) -> None:
        source = self.root / "Cliente Alfa" / "portale.txt"
        source.write_text(
            "Titolo: Portale clienti\n"
            "URL: https://example.test\n"
            "Utente: mario.rossi\n"
            "Password: Segreto-Non-Esportare-123\n",
            encoding="utf-8",
        )

        result = analyze_files(self.root, ["Cliente Alfa/portale.txt"])

        self.assertEqual(result.analyzed_files, 1)
        self.assertEqual(result.candidate_count, 1)
        candidate = result.candidates[0]
        self.assertEqual(candidate.status, "ready")
        self.assertEqual(candidate.title, "Portale clienti")
        self.assertEqual(candidate.username, "mario.rossi")
        self.assertEqual(candidate.secret_mask, "********")
        self.assertEqual(candidate.secret_length, len("Segreto-Non-Esportare-123"))
        serialized = json.dumps(asdict(result), ensure_ascii=False)
        self.assertNotIn("Segreto-Non-Esportare-123", serialized)

    def test_incomplete_candidate_is_kept_for_review(self) -> None:
        source = self.root / "Cliente Alfa" / "incompleto.env"
        source.write_text(
            "TITLE=Router\nUSERNAME=admin\nURL=https://router.test\nPASSWORD=\n",
            encoding="utf-8",
        )

        result = analyze_files(self.root, ["Cliente Alfa/incompleto.env"])

        self.assertEqual(result.candidate_count, 1)
        self.assertEqual(result.incomplete_count, 1)
        self.assertFalse(result.candidates[0].secret_present)
        self.assertEqual(result.candidates[0].secret_mask, "")

    def test_prefixed_environment_keys_are_recognized(self) -> None:
        source = self.root / "Cliente Alfa" / ".env"
        source.write_text(
            "DB_USERNAME=db-user\nDB_PASSWORD=db-secret\nDB_HOST=db.internal.test\n",
            encoding="utf-8",
        )

        result = analyze_files(self.root, ["Cliente Alfa/.env"])

        self.assertEqual(result.candidate_count, 1)
        candidate = result.candidates[0]
        self.assertEqual(candidate.username, "db-user")
        self.assertEqual(candidate.uri, "db.internal.test")
        self.assertTrue(candidate.secret_present)
        self.assertNotIn("db-secret", json.dumps(asdict(result)))

    def test_ip_address_label_populates_uri(self) -> None:
        source = self.root / "Cliente Alfa" / "router.txt"
        source.write_text(
            "Titolo: Router sede\n"
            "Username: admin\n"
            "Password: router-secret\n"
            "Indirizzo IP: 192.168.10.254\n",
            encoding="utf-8",
        )

        result = analyze_files(self.root, ["Cliente Alfa/router.txt"])

        self.assertEqual(result.candidate_count, 1)
        self.assertEqual(result.candidates[0].uri, "192.168.10.254")

    def test_embedded_ip_fills_missing_uri_but_never_reads_password_as_host(self) -> None:
        source = self.root / "Cliente Alfa" / "firewall.json"
        source.write_text(
            json.dumps(
                {
                    "titolo": "Firewall",
                    "username": "admin",
                    "password": "not-a-host-10.0.0.99",
                    "dettagli": "Gestione disponibile su 10.20.30.40 porta 443",
                }
            ),
            encoding="utf-8",
        )

        result = analyze_files(self.root, ["Cliente Alfa/firewall.json"])

        self.assertEqual(result.candidate_count, 1)
        self.assertEqual(result.candidates[0].uri, "10.20.30.40")

    def test_ipv6_address_populates_uri(self) -> None:
        source = self.root / "Cliente Beta" / "nas.env"
        source.write_text(
            "TITLE=NAS\nUSERNAME=admin\nPASSWORD=nas-secret\nIPV6=2001:db8::25\n",
            encoding="utf-8",
        )

        result = analyze_files(self.root, ["Cliente Beta/nas.env"])

        self.assertEqual(result.candidate_count, 1)
        self.assertEqual(result.candidates[0].uri, "2001:db8::25")

    def test_csv_and_nested_json_records(self) -> None:
        csv_source = self.root / "Cliente Alfa" / "credenziali.csv"
        csv_source.write_text(
            "nome,username,password,url\n"
            "CRM,crm-user,crm-secret,https://crm.test\n"
            "NAS,nas-user,nas-secret,https://nas.test\n",
            encoding="utf-8",
        )
        json_source = self.root / "Cliente Beta" / "accessi.json"
        json_source.write_text(
            json.dumps(
                {
                    "gruppo": {
                        "elementi": [
                            {
                                "service": "VPN",
                                "login": "vpn-user",
                                "pass": "vpn-secret",
                                "host": "vpn.test",
                            }
                        ]
                    }
                }
            ),
            encoding="utf-8",
        )

        result = analyze_files(
            self.root,
            ["Cliente Alfa/credenziali.csv", "Cliente Beta/accessi.json"],
        )

        self.assertEqual(result.candidate_count, 3)
        self.assertEqual({item.title for item in result.candidates}, {"CRM", "NAS", "VPN"})
        serialized = json.dumps(asdict(result))
        self.assertNotIn("crm-secret", serialized)
        self.assertNotIn("vpn-secret", serialized)

    def test_xlsx_docx_and_pdf_are_read_locally(self) -> None:
        from docx import Document
        from openpyxl import Workbook
        from reportlab.pdfgen import canvas

        xlsx_path = self.root / "Cliente Alfa" / "foglio.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Accessi"
        sheet.append(["Titolo", "Utente", "Password", "URL"])
        sheet.append(["Gestionale", "gest-user", "gest-secret", "https://gest.test"])
        workbook.save(xlsx_path)
        workbook.close()

        docx_path = self.root / "Cliente Beta" / "documento.docx"
        document = Document()
        table = document.add_table(rows=1, cols=4)
        for cell, value in zip(
            table.rows[0].cells,
            ["Nome", "Login", "Password", "Sito"],
        ):
            cell.text = value
        cells = table.add_row().cells
        for cell, value in zip(
            cells,
            ["Extranet", "extra-user", "extra-secret", "https://extra.test"],
        ):
            cell.text = value
        document.save(docx_path)

        pdf_path = self.root / "Cliente Beta" / "scheda.pdf"
        pdf = canvas.Canvas(str(pdf_path))
        pdf.drawString(72, 760, "Titolo: Firewall")
        pdf.drawString(72, 740, "Username: fw-user")
        pdf.drawString(72, 720, "Password: fw-secret")
        pdf.drawString(72, 700, "Host: firewall.test")
        pdf.save()

        result = analyze_files(
            self.root,
            [
                "Cliente Alfa/foglio.xlsx",
                "Cliente Beta/documento.docx",
                "Cliente Beta/scheda.pdf",
            ],
        )

        self.assertEqual(result.analyzed_files, 3)
        self.assertEqual(result.candidate_count, 3)
        self.assertEqual(
            {item.title for item in result.candidates},
            {"Gestionale", "Extranet", "Firewall"},
        )
        serialized = json.dumps(asdict(result))
        for secret in ("gest-secret", "extra-secret", "fw-secret"):
            self.assertNotIn(secret, serialized)

    def test_password_protected_xlsx_prompts_rejects_and_decrypts_in_memory(self) -> None:
        path = self.root / "Cliente Alfa" / "protetto.xlsx"
        document_password = "Password-file-Excel-42"
        credential_password = "Segreto-nel-foglio-99"
        relative_path = "Cliente Alfa/protetto.xlsx"
        write_encrypted_xlsx(path, document_password, credential_password)

        locked = analyze_files(self.root, [relative_path])
        self.assertEqual(locked.analyzed_files, 0)
        self.assertEqual(locked.candidate_count, 0)
        self.assertEqual(len(locked.protected_excel_issues), 1)
        self.assertEqual(locked.protected_excel_issues[0].status, "password_required")

        rejected = analyze_files(
            self.root,
            [relative_path],
            file_passwords={relative_path: "password-sbagliata"},
        )
        self.assertEqual(rejected.protected_excel_issues[0].status, "password_rejected")

        unlocked = analyze_files(
            self.root,
            [relative_path],
            file_passwords={relative_path: document_password},
        )
        self.assertEqual(unlocked.analyzed_files, 1)
        self.assertEqual(unlocked.candidate_count, 1)
        self.assertEqual(unlocked.protected_excel_issues, [])
        self.assertTrue(unlocked.candidates[0].source_password_required)
        self.assertEqual(unlocked.candidates[0].uri, "10.30.40.50")
        serialized = json.dumps(asdict(unlocked), ensure_ascii=False)
        self.assertNotIn(document_password, serialized)
        self.assertNotIn(credential_password, serialized)
        self.assertEqual([item.name for item in path.parent.iterdir()], ["protetto.xlsx"])

    def test_secure_review_cli_never_echoes_excel_passwords(self) -> None:
        path = self.root / "Cliente Beta" / "protetto.xlsx"
        document_password = "Password-documento-non-serializzare"
        credential_password = "Password-credenziale-non-serializzare"
        relative_path = "Cliente Beta/protetto.xlsx"
        write_encrypted_xlsx(path, document_password, credential_password)
        request = {
            "files": [relative_path],
            "file_passwords": [
                {"relative_path": relative_path, "password": document_password}
            ],
        }

        completed = subprocess.run(
            [
                sys.executable,
                str(Path(__file__).with_name("passbolt_review.py")),
                "--secure-json",
                "--root",
                str(self.root),
            ],
            input=json.dumps(request).encode("utf-8"),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=30,
            check=False,
        )

        self.assertEqual(completed.returncode, 0, completed.stderr.decode("utf-8"))
        envelope = json.loads(completed.stdout)
        self.assertTrue(envelope["ok"])
        self.assertEqual(envelope["result"]["candidate_count"], 1)
        process_output = (completed.stdout + completed.stderr).decode("utf-8")
        self.assertNotIn(document_password, process_output)
        self.assertNotIn(credential_password, process_output)

    def test_path_traversal_and_legacy_xls_become_safe_warnings(self) -> None:
        outside = self.root.parent / "outside-review.txt"
        outside.write_text("password=outside-secret", encoding="utf-8")
        legacy = self.root / "Cliente Alfa" / "legacy.xls"
        legacy.write_bytes(b"not-opened")
        try:
            result = analyze_files(
                self.root,
                ["../outside-review.txt", "Cliente Alfa/legacy.xls"],
            )
        finally:
            outside.unlink(missing_ok=True)

        self.assertEqual(result.analyzed_files, 0)
        self.assertEqual(result.candidate_count, 0)
        self.assertEqual(len(result.warnings), 2)
        self.assertFalse(any("outside-secret" in warning for warning in result.warnings))

    def test_file_size_limit_and_unlimited_file_selection(self) -> None:
        source = self.root / "Cliente Alfa" / "grande.txt"
        source.write_text("password=secret", encoding="utf-8")

        with mock.patch("passbolt_review.MAX_FILE_BYTES", 1):
            result = analyze_files(self.root, ["Cliente Alfa/grande.txt"])
        self.assertEqual(result.analyzed_files, 0)
        self.assertIn("limite", result.warnings[0])

        selected_files = []
        for index in range(60):
            relative_path = f"Cliente Alfa/accesso-{index}.txt"
            selected_files.append(relative_path)
            (self.root / relative_path).write_text(
                f"Titolo: Accesso {index}\n"
                f"Utente: utente-{index}\n"
                f"Password: segreto-{index}\n",
                encoding="utf-8",
            )
        result = analyze_files(self.root, selected_files)
        self.assertEqual(result.analyzed_files, 60)
        self.assertEqual(result.candidate_count, 60)

    def test_candidate_collection_exceeds_previous_aggregate_cap(self) -> None:
        source = self.root / "Cliente Alfa" / "lotto-esteso.csv"
        rows = ["nome,username,password,url"]
        rows.extend(
            f"Accesso {index},utente-{index},segreto-{index},https://host-{index}.test"
            for index in range(2_001)
        )
        source.write_text("\n".join(rows) + "\n", encoding="utf-8")

        result = analyze_files(self.root, ["Cliente Alfa/lotto-esteso.csv"])

        self.assertEqual(result.analyzed_files, 1)
        self.assertEqual(result.candidate_count, 2_001)
        self.assertEqual(len(result.candidates), 2_001)

    def test_xml_dtd_is_rejected_without_entity_expansion(self) -> None:
        source = self.root / "Cliente Alfa" / "malevolo.xml"
        source.write_text(
            '<!DOCTYPE data [<!ENTITY secret "do-not-expand">]>'
            "<data><password>&secret;</password></data>",
            encoding="utf-8",
        )

        result = analyze_files(self.root, ["Cliente Alfa/malevolo.xml"])

        self.assertEqual(result.analyzed_files, 0)
        self.assertEqual(result.candidate_count, 0)
        self.assertIn("DTD", result.warnings[0])


if __name__ == "__main__":
    unittest.main()
